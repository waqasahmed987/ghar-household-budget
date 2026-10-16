from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


GREEN = "176B55"
PALE_GREEN = "EAF4EF"
WHITE = "FFFFFF"
GRAY = "F3F5F4"


def _style_sheet(ws, widths=None):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(vertical="center")

    ws.row_dimensions[1].height = 26

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=PALE_GREEN if cell.row % 2 == 0 else WHITE)
            cell.alignment = Alignment(vertical="center")

    if widths:
        for column, width in widths.items():
            ws.column_dimensions[column].width = width
    else:
        for index, column_cells in enumerate(ws.columns, start=1):
            max_length = max(
                (len(str(cell.value or "")) for cell in column_cells),
                default=10,
            )
            ws.column_dimensions[get_column_letter(index)].width = min(max_length + 3, 34)


def create_budget_template():
    """Return a formatted Rawalpindi/Islamabad household budget workbook."""
    wb = Workbook()

    # Start Here
    ws = wb.active
    ws.title = "Start Here"
    ws.append(["Household Budget Planner"])
    ws.append(["Currency", "Pakistani Rupee (PKR)"])
    ws.append(["How to use", "Enter income, expenses, goals, and bank balances in their tabs."])
    ws.append(["Monthly actual savings", "Actual income minus actual expenses."])
    ws.append(["Daily bank updates", "Record deposits and withdrawals in Daily Bank Log."])
    ws.append(["Five-year estimate", "Change the assumptions in 5-Year Plan."])
    ws.append(["Local category examples", "Rent, groceries, utilities, school, transport, health, domestic help."])
    _style_sheet(ws, {"A": 28, "B": 100})

    # Income
    ws = wb.create_sheet("Income")
    ws.append(["Source", "Expected PKR", "Actual PKR"])
    for _ in range(5):
        ws.append(["", 0, 0])
    ws.append(["TOTAL", "=SUM(B2:B6)", "=SUM(C2:C6)"])
    _style_sheet(ws, {"A": 30, "B": 20, "C": 20})
    for row in range(2, 8):
        ws[f"B{row}"].number_format = '#,##0'
        ws[f"C{row}"].number_format = '#,##0'

    # Expenses
    expense_categories = [
        "Rent / mortgage", "Groceries", "Electricity", "Gas", "Water",
        "Internet / phone", "School fees", "Transport / fuel",
        "Healthcare / medicines", "Household supplies", "Dining out",
        "Clothing", "Childcare", "Domestic help", "Maintenance",
        "Subscriptions", "Family support", "Personal care",
        "Savings / investments", "Other",
    ]
    ws = wb.create_sheet("Expenses")
    ws.append(["Category", "Description", "Budget PKR", "Actual PKR"])
    for category in expense_categories:
        ws.append([category, "", 0, 0])
    ws.append(["TOTAL", "", "=SUM(C2:C21)", "=SUM(D2:D21)"])
    _style_sheet(ws, {"A": 30, "B": 36, "C": 20, "D": 20})
    for row in range(2, 23):
        ws[f"C{row}"].number_format = '#,##0'
        ws[f"D{row}"].number_format = '#,##0'

    # Monthly Budget
    ws = wb.create_sheet("Monthly Budget")
    ws.append(["Measure", "Planned PKR", "Actual PKR"])
    ws.append(["Income", "=Income!B7", "=Income!C7"])
    ws.append(["Expenses", "=Expenses!C22", "=Expenses!D22"])
    ws.append(["Savings", "=B2-B3", "=C2-C3"])
    ws.append([])
    ws.append(["Savings goal", "Monthly target PKR"])
    ws.append(["Emergency fund", 0])
    ws.append(["Education", 0])
    ws.append(["Home / major purchase", 0])
    ws.append(["Other goal", 0])
    ws.append(["Total goal targets", "=SUM(B7:B10)"])
    _style_sheet(ws, {"A": 32, "B": 22, "C": 22})
    for row in range(2, 5):
        ws[f"B{row}"].number_format = '#,##0'
        ws[f"C{row}"].number_format = '#,##0'
    for row in range(7, 12):
        ws[f"B{row}"].number_format = '#,##0'

    # Bank Accounts
    ws = wb.create_sheet("Bank Accounts")
    ws.append(
        [
            "Bank / account name",
            "Opening balance PKR",
            "Current balance PKR",
            "Monthly deposits PKR",
            "Monthly withdrawals PKR",
        ]
    )
    ws.append(["Account 1", 0, 0, 0, 0])
    ws.append(["Account 2", 0, 0, 0, 0])
    ws.append(["TOTAL", "=SUM(B2:B3)", "=SUM(C2:C3)", "=SUM(D2:D3)", "=SUM(E2:E3)"])
    _style_sheet(ws, {"A": 28, "B": 22, "C": 22, "D": 22, "E": 24})
    for row in range(2, 5):
        for column in "BCDE":
            ws[f"{column}{row}"].number_format = '#,##0'

    # Daily Bank Log
    ws = wb.create_sheet("Daily Bank Log")
    ws.append(["Date", "Account", "Opening PKR", "Deposit PKR", "Withdrawal PKR", "Closing PKR", "Note"])
    for row in range(2, 33):
        ws.append(["", "", 0, 0, 0, f"=C{row}+D{row}-E{row}", ""])
    _style_sheet(ws, {"A": 16, "B": 24, "C": 20, "D": 20, "E": 22, "F": 20, "G": 40})
    for row in range(2, 33):
        for column in "CDEF":
            ws[f"{column}{row}"].number_format = '#,##0'

    # Five-year plan
    ws = wb.create_sheet("5-Year Plan")
    ws.append(["Input / year", "Value"])
    ws.append(["Starting bank balance PKR", "='Bank Accounts'!C4"])
    ws.append(["Monthly saving contribution PKR", "='Monthly Budget'!C4"])
    ws.append(["Assumed annual return %", 0])
    ws.append([])
    ws.append(["Year", "Estimated balance PKR", "Total contributions PKR"])
    for year in range(1, 6):
        row = 6 + year
        if year == 1:
            formula = f"=($B$2+$B$3*12)*(1+$B$4/100)"
        else:
            prev = row - 1
            formula = f"=(B{prev}+$B$3*12)*(1+$B$4/100)"
        ws.append([year, formula, f"=$B$3*12*A{row}"])
    _style_sheet(ws, {"A": 38, "B": 26, "C": 28})
    for row in range(2, 12):
        ws[f"B{row}"].number_format = '#,##0'
    ws["B4"].number_format = "0.0"

    return wb


def template_bytes():
    buffer = BytesIO()
    create_budget_template().save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def read_uploaded_workbook(file):
    """Read known sheets from an uploaded Excel workbook into DataFrames."""
    book = pd.ExcelFile(file)
    result = {}

    for sheet_name in book.sheet_names:
        normalized = sheet_name.strip().lower()
        if normalized in {"income", "expenses", "monthly budget", "bank accounts"}:
            result[normalized] = pd.read_excel(book, sheet_name=sheet_name)

    return result
